"""엑셀 내려받기 공통. 구 서버의 표 스타일(진한 파랑 머리글, 고정 머리글, 자동 필터)을 그대로."""
from __future__ import annotations

import io
from collections.abc import Iterable, Sequence
from datetime import date, datetime
from urllib.parse import quote

from fastapi.responses import StreamingResponse

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEAD_FILL = "1F4E78"


def workbook():
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    return wb


def fill_sheet(
    ws,
    head: Sequence[str],
    rows: Iterable[Sequence],
    widths: Sequence[int] | None = None,
    wrap_cols: Sequence[int] = (),
    freeze: str = "A2",
) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws.append(list(head))
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEAD_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    n = 0
    for r in rows:
        ws.append([_cell(v) for v in r])
        n += 1
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    if wrap_cols and n:
        for row in ws.iter_rows(min_row=2):
            for i in wrap_cols:
                if i < len(row):
                    row[i].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = freeze
    if n:
        ws.auto_filter.ref = ws.dimensions


def _cell(v):
    if isinstance(v, datetime):
        return v.replace(tzinfo=None)
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, bool):
        return "O" if v else "X"
    return v


def sheet_title(name: str) -> str:
    bad = '[]:*?/\\'
    t = "".join("_" if ch in bad else ch for ch in name).strip() or "Sheet"
    return t[:31]


def to_response(wb, filename: str) -> StreamingResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    ascii_name = "".join(ch if ch.isascii() and ch.isalnum() or ch in "-_." else "_" for ch in filename)
    headers = {
        "Content-Disposition": f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(filename)}"
    }
    return StreamingResponse(buf, media_type=XLSX_MIME, headers=headers)
