"""서비스 대응일지 엑셀의 '매장정보' 시트로 매장 시스템 구성을 채운다 (2026-10-08).

엑셀 열: D 브랜드명 · E 매장명 · P 프로그램 (WINDOW / ANDROID APP / PLC).
매장명이 같은 서버 매장에 넣는다(띄어쓰기는 무시). 같은 이름의 매장이 여럿이면
브랜드가 같은 매장을 고르고, 그래도 여럿이면 건너뛴다. 매장명에 취소선이 있는 행도
넣되 미리 보기에 표시한다(--skip-struck 로 뺄 수 있다).

사용 (서버 계정으로, DB 백업 후)
  cd /opt/ddeck/backend
  sudo -u ddeck env DEBUG=false .venv/bin/python scripts/set_store_systems.py \\
      --xlsx "/경로/26년 (하남)서비스 대응일지.xlsx"            # 미리 보기
  ... --apply                                                   # 반영
"""

from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook
from sqlalchemy import select

from app.core.database import SessionLocal
from app.models.admin import CodeItem
from app.models.store import Store

PROGRAM = {
    "WINDOW": "WINDOWS",
    "WINDOWS": "WINDOWS",
    "ANDROID APP": "ANDROID",
    "ANDROID": "ANDROID",
    "PLC": "PLC",
}
LABEL = {"WINDOWS": "Windows", "ANDROID": "Android", "PLC": "PLC"}


def key(text) -> str:
    return "".join(str(text or "").split())


def read_sheet(path: Path) -> list[dict]:
    sheet = load_workbook(path, data_only=True)["매장정보"]
    rows = []
    for r in range(5, sheet.max_row + 1):
        name = sheet[f"E{r}"].value
        if not key(name):
            continue
        font = sheet[f"E{r}"].font
        raw = str(sheet[f"P{r}"].value or "").strip().upper()
        rows.append(
            {
                "row": r,
                "brand": str(sheet[f"D{r}"].value or "").strip(),
                "name": str(name).strip(),
                "program": raw,
                "system": PROGRAM.get(" ".join(raw.split())),
                "struck": bool(font and font.strike),
            }
        )
    return rows


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", type=Path, required=True)
    ap.add_argument(
        "--skip-struck", action="store_true", help="취소선 행은 넣지 않는다"
    )
    ap.add_argument(
        "--apply", action="store_true", help="실제로 반영한다 (없으면 미리 보기)"
    )
    args = ap.parse_args()
    rows = read_sheet(args.xlsx)

    db = SessionLocal()
    try:
        brands = dict(db.execute(select(CodeItem.id, CodeItem.name)).all())
        by_name = defaultdict(list)
        for store in db.scalars(select(Store).where(Store.deleted_at.is_(None))).all():
            by_name[key(store.name)].append(store)

        planned: dict = {}  # store id -> (store, system, row)
        problems = []
        for row in rows:
            where = f"{row['row']}행 {row['brand']} {row['name']}"
            if args.skip_struck and row["struck"]:
                problems.append(f"취소선이라 건너뜀: {where}")
                continue
            if row["system"] is None:
                problems.append(
                    f"프로그램 값 없음/알 수 없음({row['program'] or '-'}): {where}"
                )
                continue
            candidates = by_name.get(key(row["name"]), [])
            if len(candidates) > 1:
                same_brand = [
                    s
                    for s in candidates
                    if key(brands.get(s.brand_id)) == key(row["brand"])
                ]
                candidates = same_brand if len(same_brand) == 1 else candidates
            if not candidates:
                problems.append(f"서버에 없는 매장: {where}")
                continue
            if len(candidates) > 1:
                problems.append(f"같은 이름 매장이 여럿이라 건너뜀: {where}")
                continue
            store = candidates[0]
            earlier = planned.get(store.id)
            if earlier and earlier[1] != row["system"]:
                problems.append(
                    f"엑셀에 값이 서로 다름({earlier[2]['row']}행 {LABEL[earlier[1]]} / "
                    f"{row['row']}행 {LABEL[row['system']]}), 건너뜀: {store.name}"
                )
                planned[store.id] = (store, None, row)
                continue
            if not earlier:
                planned[store.id] = (store, row["system"], row)

        changes = [
            (store, system, row)
            for store, system, row in planned.values()
            if system and store.system_type != system
        ]
        same = sum(
            1 for s, system, _ in planned.values() if system and s.system_type == system
        )
        print(f"엑셀 매장 {len(rows)}행 · 서버 매장과 맞춘 {len(planned)}곳")
        print(f"\n[바꿀 매장] {len(changes)}곳 (이미 같은 값 {same}곳)")
        for store, system, row in sorted(changes, key=lambda c: c[2]["row"]):
            before = LABEL.get(store.system_type, "미지정")
            mark = " (엑셀 취소선)" if row["struck"] else ""
            print(
                f"  {row['row']:>4}행 {store.name}: {before} -> {LABEL[system]}{mark}"
            )
        print(f"\n[확인 필요] {len(problems)}건")
        for line in problems:
            print("  " + line)
        if args.apply:
            for store, system, _ in changes:
                store.system_type = system
            db.commit()
            print("\n반영했습니다.")
        else:
            db.rollback()
            print("\n미리 보기입니다. 반영하려면 --apply 를 붙이세요.")
    finally:
        db.close()


if __name__ == "__main__":
    main()
