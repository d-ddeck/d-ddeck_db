"""Extensible service work types: API, history, filters and exports in a temporary DB."""

from __future__ import annotations

import io
import os
import sys
import tempfile
from pathlib import Path
from unittest.mock import patch
from uuid import UUID, uuid4

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
TEMP = tempfile.TemporaryDirectory(prefix="ddeck-work-types-")
os.environ.update(
    DATABASE_URL=f"sqlite+pysqlite:///{TEMP.name}/test.db",
    ENVIRONMENT="test",
    DEBUG="false",
    AUTH_RATE_LIMIT_ENABLED="false",
    SCHEDULER_ENABLED="false",
    STORAGE_DIR=f"{TEMP.name}/storage",
    FIRST_SUPERADMIN_EMAIL="admin@ddeck.local",
    FIRST_SUPERADMIN_PASSWORD="admin1234",
)
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.api.v1 import service
from app.core.database import SessionLocal, engine
from app.main import app
from app.models.admin import AuditLog
from app.models.service import ServiceTicket, ServiceTicketNumber

try:
    with TestClient(app) as client:

        def login(password):
            return {
                "Authorization": "Bearer "
                + client.post(
                    "/api/v1/auth/login", json={"email": "admin", "password": password}
                ).json()["access_token"]
            }

        headers = login("admin1234")

        def call(method, path, body=None, expected=200):
            r = client.request(method, "/api/v1" + path, json=body, headers=headers)
            assert r.status_code == expected, (path, r.status_code, r.text)
            return r.json()

        call(
            "POST",
            "/auth/change-password",
            {"current_password": "admin1234", "new_password": "admin5678"},
        )
        headers = login("admin5678")
        group = call("GET", "/admin/codes/SERVICE_WORK_TYPE")
        types = {i["code"]: i["id"] for i in group["items"]}
        assert set(types) == {"AS", "CS", "PO"}
        custom = call(
            "POST",
            f"/admin/codes/{group['id']}/items",
            {"code": "INSTALL", "name": "설치 지원"},
            201,
        )
        types["INSTALL"] = custom["id"]
        category = next(
            i["id"]
            for i in call("GET", "/admin/codes/SERVICE_CATEGORY")["items"]
            if i["code"] == "PROGRAM"
        )
        base = {
            "title": "업무 구분 검사",
            "category_id": category,
            "received_at": "2026-09-27T00:00:00Z",
        }

        def create(**extra):
            return call("POST", "/service/tickets", {**base, **extra}, 201)

        legacy = create()
        assert legacy["work_type_id"] is None and legacy["work_type"] is None
        assert legacy["ticket_no"].startswith("AS-")
        tickets = {
            code: create(work_type_id=identifier) for code, identifier in types.items()
        }
        for code, ticket in tickets.items():
            assert ticket["work_type"]["code"] == code
            assert ticket["ticket_no"].startswith(code + "-")
            assert ticket["ticket_no"].endswith("-0002" if code == "AS" else "-0001")
            query = f"work_type_id={types[code]}"
            assert call("GET", "/service/tickets?" + query)["total"] == 1
            assert call("GET", "/service/stats/summary?" + query)["total"] == 1
            grouped = call("GET", "/service/stats/grouped?group_by=work_type&" + query)
            assert grouped["total"] == 1, grouped
            assert len(grouped["buckets"]) == 1 and grouped["buckets"][0]["count"] == 1
            assert grouped["buckets"][0]["key"] == types[code]
            trend = call("GET", "/service/stats/trend?interval=month&" + query)
            assert sum(p["received"] for p in trend["points"]) == 1
            cross = call(
                "GET", "/service/stats/crosstab?rows=year&cols=category&" + query
            )
            assert cross["total_tickets"] == 1
        assert call("GET", "/service/tickets?missing=work_type")["total"] == 1
        assert call("GET", "/service/stats/summary?missing=work_type")["total"] == 1
        assert call("GET", "/service/tickets?q=INSTALL")["total"] == 1
        assert call("GET", "/service/tickets?q=설치 지원")["total"] == 1
        for invalid in (category, str(uuid4())):
            call("POST", "/service/tickets", {**base, "work_type_id": invalid}, 400)
        ticket_path = "/service/tickets/" + tickets["CS"]["id"]
        changed = call("PATCH", ticket_path, {"work_type_id": types["PO"]})
        assert changed["work_type"]["code"] == "PO"
        assert changed["ticket_no"].startswith("PO-")
        assert changed["ticket_no"].endswith("-0002")
        assert changed["ticket_no"] != tickets["CS"]["ticket_no"]
        assert (
            call("PATCH", ticket_path, {"work_type_id": types["PO"]})["ticket_no"]
            == changed["ticket_no"]
        )
        assert (
            call("PATCH", ticket_path, {"title": "동일 번호 유지"})["ticket_no"]
            == changed["ticket_no"]
        )
        assert any(
            tickets["CS"]["ticket_no"] in (log["content"] or "")
            and changed["ticket_no"] in (log["content"] or "")
            for log in changed["logs"]
        )
        assert (
            call("PATCH", ticket_path, {"title": "구매로 변경"})["work_type_id"]
            == types["PO"]
        )
        cleared = call("PATCH", ticket_path, {"work_type_id": None})
        assert cleared["work_type"] is None
        assert cleared["ticket_no"].startswith("AS-")
        assert call("GET", "/service/tickets?missing=work_type")["total"] == 2
        item_path = "/admin/codes/items/" + custom["id"]
        call("PATCH", item_path, {"name": "현장 설치", "is_active": False})
        custom_path = "/service/tickets/" + tickets["INSTALL"]["id"]
        assert call("GET", custom_path)["work_type"]["name"] == "현장 설치"
        assert call("GET", item_path + "/usage")["count"] >= 1
        call("POST", "/service/tickets", {**base, "work_type_id": custom["id"]}, 400)
        call("PATCH", ticket_path, {"work_type_id": custom["id"]}, 400)
        call(
            "PATCH",
            custom_path,
            {"work_type_id": custom["id"], "title": "기존 사용 중지 항목 유지"},
        )
        call("DELETE", item_path)
        assert call("GET", custom_path)["work_type"]["code"] == "INSTALL"
        call(
            "PATCH",
            custom_path,
            {"work_type_id": custom["id"], "title": "삭제 후에도 이력 유지"},
        )
        call("POST", "/service/tickets", {**base, "work_type_id": custom["id"]}, 400)
        r = client.get(
            "/api/v1/service/tickets/export.xlsx",
            params={"work_type_id": types["PO"]},
            headers=headers,
        )
        assert r.status_code == 200, r.text
        rows = list(load_workbook(io.BytesIO(r.content)).active.values)
        assert rows[0][-2:] == ("업무 구분 코드", "업무 구분"), rows[0]
        assert len(rows) == 2 and rows[1][-2:] == ("PO", "구매"), rows
        r = client.get(
            "/api/v1/service/stats/all.xlsx",
            params={"work_type_id": types["PO"]},
            headers=headers,
        )
        assert r.status_code == 200
        workbook = load_workbook(io.BytesIO(r.content))
        assert "업무 구분" in workbook.sheetnames
        assert list(workbook["업무 구분"].values)[1][:2] == ("구매", 1)
        # Reissued numbers remain reserved in each work type sequence.
        assert create(work_type_id=types["CS"])["ticket_no"].endswith("-0002")
        assert create(work_type_id=types["PO"])["ticket_no"].endswith("-0003")
        call(
            "PUT",
            "/admin/settings/SERVICE",
            {
                "settings": [
                    {
                        "key": "ticket_prefix",
                        "value": "DEFAULT",
                        "value_type": "string",
                        "label": "접수번호 접두어",
                        "is_public": True,
                    }
                ]
            },
        )
        assert create()["ticket_no"].startswith("DEFAULT-")
        assert create(work_type_id=types["AS"])["ticket_no"].startswith("AS-")
        # Custom codes may include SQL LIKE wildcards or use all 60 characters.
        for code in ("XAY", "X_Y", "X%Y", "LONG" * 15):
            item = call(
                "POST",
                f"/admin/codes/{group['id']}/items",
                {"code": code, "name": code},
                201,
            )
            first = create(work_type_id=item["id"])
            second = create(work_type_id=item["id"])
            assert first["ticket_no"].startswith(code + "-")
            assert first["ticket_no"].endswith("-0001")
            assert second["ticket_no"].endswith("-0002")
        # Imported records keep their legacy ID, but display/export the new number.
        with SessionLocal() as db:
            imported = db.get(ServiceTicket, UUID(legacy["id"]))
            imported.legacy_no = 42
            imported.ticket_no = "42"
            db.commit()
        legacy_path = "/service/tickets/" + legacy["id"]
        assert call("GET", legacy_path)["legacy_no"] == 42
        converted = call("PATCH", legacy_path, {"work_type_id": types["PO"]})
        assert converted["ticket_no"].startswith("PO-")
        assert converted["legacy_no"] is None and converted["id"] == legacy["id"]
        listed = call("GET", "/service/tickets?q=" + converted["ticket_no"])
        assert listed["items"][0]["legacy_no"] is None
        assert listed["items"][0]["ticket_no"] == converted["ticket_no"]
        assert f"접수번호 재발급: 42 → {converted['ticket_no']}" in converted["notices"]
        saved_again = call("PATCH", legacy_path, {"work_type_id": types["PO"]})
        assert saved_again["ticket_no"] == converted["ticket_no"]
        assert not any("접수번호 재발급" in notice for notice in saved_again["notices"])
        with SessionLocal() as db:
            assert db.get(ServiceTicket, UUID(legacy["id"])).legacy_no == 42
            assert db.get(ServiceTicketNumber, "42") is not None
            audits = db.scalars(
                select(AuditLog).where(AuditLog.entity_id == legacy["id"])
            )
            assert any(
                (row.changes or {}).get("ticket_no") == ["42", converted["ticket_no"]]
                for row in audits
            )
        exported = client.get("/api/v1/service/tickets/export.xlsx", headers=headers)
        exported_rows = list(load_workbook(io.BytesIO(exported.content)).active.values)
        assert any(row[0] == converted["ticket_no"] for row in exported_rows[1:])

        # A uniqueness collision must roll back only the savepoint, then retry.
        original_allocator = service._next_ticket_no

        def collide_once(db, attempt=0, *, received_at, prefix=None):
            if attempt == 0:
                return tickets["PO"]["ticket_no"]
            return original_allocator(db, attempt, received_at=received_at, prefix=prefix)

        with patch.object(service, "_next_ticket_no", side_effect=collide_once):
            retried = call("PATCH", legacy_path, {"work_type_id": types["CS"]})
        assert retried["ticket_no"].startswith("CS-")
        assert retried["work_type_id"] == types["CS"]
        with patch.object(
            service, "_next_ticket_no", return_value=tickets["PO"]["ticket_no"]
        ):
            call("PATCH", legacy_path, {"work_type_id": types["PO"]}, 409)
        unchanged = call("GET", legacy_path)
        assert unchanged["ticket_no"] == retried["ticket_no"]
        assert unchanged["work_type_id"] == types["CS"]
        # Validation after allocation must roll back the number and work type too.
        with SessionLocal() as db:
            reserved = set(db.scalars(select(ServiceTicketNumber.ticket_no)))
        call(
            "PATCH", legacy_path, {"work_type_id": types["PO"], "is_rental": True}, 400
        )
        assert call("GET", legacy_path)["ticket_no"] == retried["ticket_no"]
        with SessionLocal() as db:
            assert set(db.scalars(select(ServiceTicketNumber.ticket_no))) == reserved
        historical = create(work_type_id=types["AS"], received_at="2020-01-31T15:00:00Z")
        assert historical["ticket_no"] == "AS-202002-0001"
        historical_path = "/service/tickets/" + historical["id"]
        changed = call("PATCH", historical_path, {"work_type_id": types["PO"]})
        assert changed["ticket_no"] == "PO-202002-0001"
        same_month = call("PATCH", historical_path, {"received_at": "2020-02-15T00:00:00Z"})
        assert same_month["ticket_no"] == changed["ticket_no"]
        new_month = call("PATCH", historical_path, {"received_at": "2020-01-31T14:59:59Z"})
        assert new_month["ticket_no"] == "PO-202001-0001"
        restored_month = call("PATCH", historical_path, {"received_at": "2020-02-01T00:00:00Z"})
        assert restored_month["ticket_no"] == "PO-202002-0002"
        both = call("PATCH", historical_path, {"work_type_id": types["CS"], "received_at": "2019-12-31T15:00:00Z"})
        assert both["ticket_no"] == "CS-202001-0001"
        with SessionLocal() as db:
            imported = db.get(ServiceTicket, UUID(historical["id"]))
            imported.legacy_no, imported.ticket_no = 987654, "987654"
            db.commit()
        imported_result = call("PATCH", historical_path, {"work_type_id": types["AS"]})
        assert imported_result["ticket_no"] == "AS-202001-0001"
        assert imported_result["legacy_no"] is None
        with SessionLocal() as db:
            assert db.get(ServiceTicket, UUID(historical["id"])).legacy_no == 987654
        print("PASS: occurrence month, KST boundary, month edits and imported renumbering")
        print(
            "PASS: renumbering, audit/export, reserved numbers, collision retry and rollback"
        )
        # Creation can transition atomically; invalid completion leaves no ticket.
        progress = create(initial_status="IN_PROGRESS", note="방문 전 연락")
        assert progress["status"] == "IN_PROGRESS" and progress["started_at"]
        assert any("방문 전 연락" in log["content"] for log in progress["logs"])
        with SessionLocal() as db:
            before_ids = set(db.scalars(select(ServiceTicket.id)))
        call("POST", "/service/tickets", {**base, "initial_status": "COMPLETED"}, 400)
        with SessionLocal() as db:
            assert set(db.scalars(select(ServiceTicket.id))) == before_ids
            from app.models.enums import Role, UserStatus
            from app.models.user import User
            db.add(User(email="worker@example.com", full_name="서비스 담당", password_hash="unused", status=UserStatus.APPROVED, role=Role.MEMBER))
            db.commit()
        responder = call("GET", "/admin/codes/SERVICE_RESPONDER")["items"][0]["id"]
        completed = create(initial_status="COMPLETED", result_note="수리 완료", responder_ids=[responder])
        assert completed["status"] == "COMPLETED"
        assert completed["started_at"] and completed["completed_at"]
        assert completed["result_note"] == "수리 완료"
        assert [log["to_status"] for log in completed["logs"]] == ["RECEIVED", "IN_PROGRESS", "COMPLETED"]
        call("POST", "/service/tickets", {**base, "initial_status": "CANCELED"}, 422)
        print("PASS: direct progress/completion, notes, required completion fields and rollback")
        call("DELETE", "/admin/codes/items/" + types["AS"])
        print(
            "PASS: defaults/custom codes, validation, legacy null, edit/history, filters/statistics and Excel"
        )
    # Startup must not restore deleted types or duplicate the defaults.
    with TestClient(app):
        pass
    with TestClient(app) as client:
        headers = login("admin5678")
        items = call("GET", "/admin/codes/SERVICE_WORK_TYPE")["items"]
        assert [i["code"] for i in items].count("PO") == 1
        assert not any(i["code"] == "AS" for i in items)
        assert not any(i["code"] == "INSTALL" for i in items)
    print("PASS: restart seeding is idempotent and preserves deleted custom types")
finally:
    engine.dispose()
    TEMP.cleanup()
